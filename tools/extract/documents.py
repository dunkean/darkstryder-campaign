"""Offline direct conversion: HTML, Word, workbook and text. Never run macros/scripts."""
import hashlib
import io
import json
import os
import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import unquote,urlsplit
from bs4 import BeautifulSoup,NavigableString
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image
from common import rebuild,write_json

ROOT=Path(__file__).resolve().parents[2];ENGINE='local-document-conversion-v1'
def file_hash(path):
    with path.open('rb') as f:return hashlib.file_digest(f,'sha256').hexdigest()
def table(rows):
    if not rows:return ''
    width=max(map(len,rows));escaped=[[str(c if c is not None else '').replace('|','\\|').replace('\n','<br>') for c in row]+['']*(width-len(row)) for row in rows]
    return '\n\n'+'\n'.join(['| '+' | '.join(escaped[0])+' |','| '+' | '.join(['---']*width)+' |',*['| '+' | '.join(row)+' |' for row in escaped[1:]]])+'\n\n'
def save_image(data,folder,report,origin):
    try:
        with Image.open(io.BytesIO(data)) as picture:
            picture.verify();fmt=picture.format.lower();ext={'jpeg':'jpg','tiff':'tif'}.get(fmt,fmt)
        if ext not in ['jpg','png','gif','webp','bmp','tif']:raise ValueError('Unsupported image format')
        name=hashlib.sha256(data).hexdigest()+'.'+ext;target=folder/'images'/name
        if not target.exists():target.write_bytes(data)
        report['media'].append({'file':'images/'+name,'origin':origin,'sha256':hashlib.sha256(data).hexdigest()})
        return '../images/'+name
    except Exception as error:
        report['warnings'].append({'kind':'unreadable-image','origin':origin,'error':str(error)});return None
def html_document(path,folder,report,allowed_root):
    soup=BeautifulSoup(path.read_bytes(),'html.parser');report['encoding']=soup.original_encoding
    for element in soup(['script','style','iframe','object','embed']):element.decompose()
    def convert(element):
        if isinstance(element,NavigableString):return str(element)
        name=element.name;children=lambda:''.join(convert(c) for c in element.children)
        if name=='img':
            url=element.get('src','');parts=urlsplit(url)
            if parts.scheme or parts.netloc:
                report['warnings'].append({'kind':'remote-image-not-fetched','url':url});return '\n[Image distante non téléchargée : '+url+']\n'
            target=(path.parent/unquote(parts.path).replace('\\','/')).resolve()
            if not target.is_relative_to(allowed_root) or not target.is_file():
                report['warnings'].append({'kind':'missing-local-image','url':url});return '\n[Image locale absente : '+url+']\n'
            link=save_image(target.read_bytes(),folder,report,url)
            return '\n!['+element.get('alt','Image').replace(']','')+']('+link+')\n' if link else '\n[Image non décodée]\n'
        if name=='table':
            if element.find('table') or any(len(c.get_text())>500 or c.find(['h1','h2','h3','p','div','ul','ol']) for c in element.find_all(['td','th'])):return '\n\n'+children()+'\n\n'
            return table([[convert(c) for c in row.find_all(['th','td'],recursive=False)] for row in element.find_all('tr') if row.find_parent('table')==element])
        if name=='frame':return '\nCadre local : `'+element.get('src','')+'` (document inventorié séparément).\n'
        if name in ['br','hr']:return '\n' if name=='br' else '\n\n---\n\n'
        text=children()
        if name in ['h1','h2','h3','h4','h5','h6']:return '\n\n'+'#'*int(name[1])+' '+text.strip()+'\n\n'
        if name in ['strong','b']:return '**'+text+'**'
        if name in ['em','i']:return '*'+text+'*'
        if name=='li':return '\n- '+text.strip()+'\n'
        if name=='a':
            url=element.get('href','');parts=urlsplit(url)
            if parts.scheme not in ['', 'http','https','mailto']:return text
            label=text.strip() if '![' in text else text.strip().replace(']','')
            return '['+label+']('+url.replace(' ','%20')+')' if url else text
        if name in ['p','div','section','article','tr','pre','ul','ol']:return '\n\n'+text+'\n\n'
        return text
    return re.sub(r'\n[ \t]*\n(?:[ \t]*\n)+','\n\n',convert(soup)).strip(),{'kind':'html','title':soup.title.get_text() if soup.title else '', 'warnings':report['warnings']}
def word_document(path,folder,report):
    document=Document(path);parts=[];structure=[]
    for item in document.element.body:
        if item.tag.endswith('}p'):
            paragraph=Paragraph(item,document);text=paragraph.text;style=paragraph.style.name
            if style.startswith('Heading '):text='#'*int(style.split()[-1])+' '+text
            elif style.startswith('List'):text='- '+text
            parts.append(text);structure.append({'kind':'paragraph','style':style,'text':paragraph.text})
            for blip in item.xpath('.//a:blip'):
                rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if rid and rid in document.part.related_parts:
                    link=save_image(document.part.related_parts[rid].blob,folder,report,rid)
                    if link:parts.append('![Image]('+link+')')
        elif item.tag.endswith('}tbl'):
            rows=[[cell.text for cell in row.cells] for row in Table(item,document).rows];parts.append(table(rows));structure.append({'kind':'table','rows':rows})
    for section in document.sections:
        for part in [section.header,section.footer]:
            texts=[p.text for p in part.paragraphs if p.text.strip()]
            if texts:parts.append('\n### En-tête / pied de page\n'+'\n'.join(texts))
    return '\n\n'.join(parts),{'kind':'docx','elements':structure}
def legacy_word(path,folder,report,runtime):
    if path.read_bytes()[:5]==b'{\\rtf':
        from striprtf.striprtf import rtf_to_text
        text=rtf_to_text(path.read_bytes().decode('cp1252'),errors='strict')
        report['warnings'].append({'kind':'rtf-layout-approximate','note':'RTF stored with .doc extension; direct text extraction, no artificial pagination.'})
        return text,{'kind':'rtf','warnings':report['warnings']}
    home=runtime/'tools/antiword/usr/share/antiword';binary=runtime/'tools/antiword/usr/bin/antiword'
    if not binary.is_file():raise RuntimeError('Local antiword required for legacy .doc')
    result=subprocess.run([str(binary),'-m','UTF-8.txt','-w','0',str(path)],env={**os.environ,'ANTIWORDHOME':str(home)},capture_output=True,check=True,timeout=60)
    text=result.stdout.decode('utf-8');report['warnings'].append({'kind':'legacy-word-layout-approximate','note':'Text and tables via antiword; headings/layout require original verification.'})
    import olefile
    pictures=[]
    with olefile.OleFileIO(path) as doc:
        for stream in doc.listdir():
            if doc.get_size(stream)>32*1024*1024:raise RuntimeError('OLE stream too large')
            data=doc.openstream(stream).read()
            for signature,end in [(b'\x89PNG\r\n\x1a\n',b'IEND\xaeB`\x82'),(b'\xff\xd8\xff',b'\xff\xd9')]:
                offset=0
                while True:
                    start=data.find(signature,offset)
                    if start<0:break
                    stop=data.find(end,start+len(signature));offset=start+len(signature)
                    if stop<0:continue
                    link=save_image(data[start:stop+len(end)],folder,report,'OLE stream '+ '/'.join(stream))
                    if link and link not in pictures:pictures.append(link)
    if pictures:text+='\n\n### Images embarquées récupérées — association au texte non établie\n\n'+'\n'.join('![Image extraite]('+p+')' for p in pictures)
    report['warnings'].append({'kind':'legacy-word-media-recovery-limited','note':'Only decodable embedded PNG/JPEG recovered; original retained for other image types.'})
    return text,{'kind':'legacy-doc','method':'antiword + verified OLE image recovery','warnings':report['warnings']}
def convert_source(source,config):
    runtime=Path(config['runtimeRoot']).resolve();base=Path(config['sourceRoot']).resolve() if source['locator']['root']=='drive' else runtime
    path=(base/source['locator']['path']).resolve();folder=(runtime/source['output']).resolve()
    if not path.is_relative_to(base) or not folder.is_relative_to(runtime):raise RuntimeError('Path confinement failed')
    if file_hash(path)!=source['sha256']:raise RuntimeError('Original changed')
    manifest_path=folder/'manifest.json'
    if manifest_path.exists():
        old=json.loads(manifest_path.read_text())
        if old.get('engineId')!=ENGINE or old.get('sha256')!=source['sha256']:raise RuntimeError('Conversion provenance mismatch')
        if (source['format']!='html' or old.get('converterVersion')==2) and len(old['pages'])==source['pages'] and all((folder/'pages'/f'{p["page"]:04}.md').exists() and (folder/'structure'/f'{p["page"]:04}.json').exists() for p in old['pages']):return False
    for name in ['pages','images','structure']:(folder/name).mkdir(parents=True,exist_ok=True)
    report={'warnings':[],'media':[]};fmt=source['format'];units=[]
    if fmt=='html':units=[html_document(path,folder,report,base)]
    elif fmt=='docx':units=[word_document(path,folder,report)]
    elif fmt=='doc':units=[legacy_word(path,folder,report,runtime)]
    elif fmt=='xlsx':
        book=load_workbook(path,read_only=False,data_only=False)
        for sheet in book.worksheets:
            cells=sorted((cell for cell in sheet._cells.values() if cell.value is not None),key=lambda c:(c.row,c.column))
            if len(cells)>1000000:raise RuntimeError('Worksheet exceeds actual cell bound')
            max_row=max((c.row for c in cells),default=1);max_col=max((c.column for c in cells),default=1)
            if max_row*max_col>1000000:
                # Preserve sparse positions without allocating a million empty styled cells.
                rows=[['Cellule','Valeur / formule'],*[[c.coordinate,c.value] for c in cells]]
                report['warnings'].append({'kind':'sparse-worksheet','sheet':sheet.title,'note':'Coordinate/value table; empty formatting-only cells omitted.'})
            else:rows=[list(row) for row in sheet.iter_rows(min_row=1,max_row=max_row,max_col=max_col,values_only=True)]
            text='# '+sheet.title+'\n\n'+table(rows)
            for image in sheet._images:
                link=save_image(image._data(),folder,report,sheet.title)
                if link:text+='\n![Image de feuille]('+link+')\n'
            units.append((text,{'kind':'worksheet','name':sheet.title,'rows':rows,'formulas':'preserved, never evaluated'}))
        book.close()
    else:
        raw=path.read_bytes()
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('cp1252');report['warnings'].append({'kind':'encoding-fallback-cp1252'})
        units=[(text,{'kind':fmt})]
    if len(units)!=source['pages']:raise RuntimeError('Unit count differs from inventory')
    pages=[]
    for number,(text,structure) in enumerate(units,1):
        flags=[]
        if len(text.strip())<80:flags.append('low-text-or-illustration')
        if '\ufffd' in text:flags.append('replacement-characters')
        header=f'<!-- source: {source["id"]}; document_unit: {number}; format: {fmt}; engine: {ENGINE} -->\n\n'
        output=folder/'pages'/f'{number:04}.md';temp=output.with_suffix('.tmp');temp.write_text(header+text+'\n',encoding='utf-8');temp.replace(output)
        write_json(folder/'structure'/f'{number:04}.json',structure)
        pages.append({'page':number,'unitLabel':source['unitLabel'],'characters':len(text),'flags':flags})
    write_json(manifest_path,{'sourceId':source['id'],'sha256':source['sha256'],'engineId':ENGINE,'converterVersion':2,'engine':'Direct local document conversion (no OCR / no remote fetching)','format':fmt,'pages':pages,**report})
    rebuild(folder,source);return True
def main():
    config=json.loads((ROOT/'config.local.json').read_text(encoding='utf-8-sig'));runtime=Path(config['runtimeRoot'])
    sys.path.insert(0,str(runtime/'tools/document-libs'))
    sources=json.loads((ROOT/'catalog/sources.json').read_text());converted=0
    import fcntl
    with (runtime/'ocr-job.lock').open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        for source in sources:
            if source.get('format','pdf')!='pdf':
                converted+=convert_source(source,config)
                print('Converted native document:',source['id'],source['name'],flush=True)
    print(json.dumps({'converted':converted}))
if __name__=='__main__':main()
